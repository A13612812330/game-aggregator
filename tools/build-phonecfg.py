#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把「基础测试数据.xlsx」（手机跑 PC 游戏的实测配置记录）转成 GameHub 站点用的 JSON。

输入：E:\新建文件夹\基础测试数据.xlsx  （Sheet1，1062 行 x 18 列）
输出：data/phonecfg.json  —— 手机端实测配置库索引

表格列语义（据实还原，非猜测）：
  A 机型芯片代号        B 游戏名
  C 是否可玩(是/否)      D 兼容层版本(proton/wine)
  E 运行模式(兼容/性能/快速)  F 驱动/GPU(turnip_v24.2.0_R22 / Adreno_814 / 8Elite-800.46)
  G DXVK 版本           H vkd3d 版本
  I 运行库(Box64/Fex)   J 帧率文本(60帧 / 30-60帧 / 10-20帧…)
  K 备注(问题/手柄设置)   L 主程序 exe 名
  M 状态(1=有配置 0=无)   N 零星补充
"""
import json
import os
import re
import time
import openpyxl

SRC = r'E:\新建文件夹\基础测试数据.xlsx'
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data')
OUT = os.path.join(OUT_DIR, 'phonecfg.json')

# 机型代号 → 可读平台名（表里是小米/高通内部代号，按实测机型还原）
CHIP = {
    '8gen1': '骁龙 8 Gen 1',
    '8e':    '骁龙 8 Elite',
    '8egn1': '骁龙 8 Gen 1',
    '8gen3': '骁龙 8 Gen 3',
    '870':   '骁龙 870',
}

def s(v):
    return '' if v is None else str(v).strip()

# ---------- 帧率解析：把自由文本归一成可排序的数值区间 ----------
def parse_fps(txt):
    """返回 (lo, hi, mid, label, tier)。tier 用于粗粒度分档。"""
    t = s(txt)
    if not t:
        return (None, None, None, '', '')
    if re.search(r'不显示', t):
        smooth = '流畅' in t and '比较' not in t
        return (None, None, 60 if smooth else 30, t, '流畅' if smooth else '可玩')
    if re.search(r'闪退|卡死|打不开|玩不了|黑屏|卡关机|无法|崩溃', t):
        return (0, 0, 0, t, '不可用')
    nums = [float(x) for x in re.findall(r'\d+(?:\.\d+)?', t)]
    if not nums:
        return (None, None, None, t, '')
    lo, hi = min(nums), max(nums)
    if hi > 144: hi = lo          # 误把 "100帧" 之类之外的大数当噪声
    mid = (lo + hi) / 2
    if hi >= 55:      tier = '流畅'
    elif hi >= 28:    tier = '可玩'
    elif hi >= 15:    tier = '勉强'
    else:             tier = '卡顿'
    return (lo, hi, round(mid, 1), t, tier)

# ---------- 兼容层归一（proton/wine 两系） ----------
def layer_family(v):
    v = s(v).lower()
    if v.startswith('proton'): return 'proton'
    if v.startswith('wine'):   return 'wine'
    return ''

def norm_key(t):
    """归一化钥匙。
    注意：中文游戏的「去非字母数字」会得到空串，绝不能用来单独做聚合键，
    所以这里保留 CJK 字符，只做「大小写/空白/标点」层面的归一。
    """
    s0 = s(t).lower()
    s0 = re.sub(r'[\s\u3000]+', '', s0)                     # 去空白
    s0 = re.sub(r'[·・:：,，.。!！?？"\'“”‘’()（）\[\]【】<>《》|｜/\\~～\-—_+*&#@$%^&;；]', '', s0)
    return s0

def title_keys(title):
    """游戏名 → 多把归一化钥匙（与 data/bannerhub.js 的 titleKeys 保持同构）。
    返回的钥匙必须非空，否则会污染索引。
    """
    t = s(title)
    if not t: return []
    out = set()
    for part in re.split(r'[/／|｜]', t):
        p = part.strip()
        if not p: continue
        k1 = norm_key(p)
        if k1: out.add(k1)
        # 去掉括号内容 / 副标题分隔符再生成一把
        base = re.split(r'[：:（(\[【]', p)[0].strip()
        k2 = norm_key(base)
        if k2: out.add(k2)
    out.discard('')
    return sorted(out)

def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb['Sheet1']

    rows = []
    for r in range(2, ws.max_row + 1):
        g = s(ws.cell(row=r, column=2).value)
        if not g:
            continue                                    # 无游戏名 → 非数据行
        chip_raw = s(ws.cell(row=r, column=1).value)
        # A 列有脏值（网址/数字/处理器…），只认已知代号
        chip = chip_raw if chip_raw in CHIP else ''
        if not chip and chip_raw:
            for k in CHIP:
                if k in chip_raw:
                    chip = k; break
        playable = s(ws.cell(row=r, column=3).value)
        fps_lo, fps_hi, fps_mid, fps_label, tier = parse_fps(ws.cell(row=r, column=10).value)
        d = {
            'chipRaw': chip_raw,
            'chip': chip,
            'chipName': CHIP.get(chip, ''),
            'game': g,
            'playable': playable,                        # 是 / 否 / ''
            'ok': playable == '是',
            'layer': s(ws.cell(row=r, column=4).value),
            'layerFam': layer_family(ws.cell(row=r, column=4).value),
            'mode': s(ws.cell(row=r, column=5).value),
            'gpu': s(ws.cell(row=r, column=6).value),
            'dxvk': s(ws.cell(row=r, column=7).value),
            'vkd3d': s(ws.cell(row=r, column=8).value),
            'runtime': s(ws.cell(row=r, column=9).value),
            'fpsLabel': fps_label,
            'fpsLo': fps_lo, 'fpsHi': fps_hi, 'fpsMid': fps_mid,
            'fpsTier': tier,
            'note': s(ws.cell(row=r, column=11).value),
            'exe': s(ws.cell(row=r, column=12).value),
            'hasCfg': s(ws.cell(row=r, column=13).value) == '1',
            'row': r,
        }
        if d['vkd3d'].lower() == 'none': d['vkd3d'] = ''   # 表里 "None" 表示不用该组件
        rows.append(d)

    # ---------- 按游戏聚合 ----------
    by_game = {}
    for d in rows:
        k = norm_key(d['game'])
        g = by_game.get(k)
        if not g:
            g = {'k': k, 'title': d['game'], 'n': 0, 'okN': 0, 'chips': set(),
                 'gpus': set(), 'tiers': set(), 'bestMid': None, 'bestLabel': '',
                 'notes': [], 'exes': set(), 'layers': set(), 'modes': set(),
                 'hasCfgN': 0, 'keys': title_keys(d['game'])}
            by_game[k] = g
        g['n'] += 1
        if d['ok']: g['okN'] += 1
        if d['chip']: g['chips'].add(d['chip'])
        if d['gpu']: g['gpus'].add(d['gpu'])
        if d['fpsTier']: g['tiers'].add(d['fpsTier'])
        if d['hasCfg']: g['hasCfgN'] += 1
        if d['fpsMid'] is not None and d['fpsMid'] > 0 and (g['bestMid'] is None or d['fpsMid'] > g['bestMid']):
            g['bestMid'] = d['fpsMid']; g['bestLabel'] = d['fpsLabel']
        if d['note']: g['notes'].append({'chip': d['chipRaw'], 'note': d['note']})
        if d['exe'].lower().endswith('.exe'): g['exes'].add(d['exe'])
        if d['layer']: g['layers'].add(d['layer'])
        if d['mode']: g['modes'].add(d['mode'])

    games = []
    for k, g in by_game.items():
        games.append({
            'k': k, 'title': g['title'], 'n': g['n'], 'okN': g['okN'],
            'chips': sorted(g['chips']), 'gpus': sorted(g['gpus'])[:6],
            'tiers': sorted(g['tiers']),
            'bestMid': g['bestMid'], 'bestLabel': g['bestLabel'],
            'notes': g['notes'][:5], 'exes': sorted(g['exes'])[:4],
            'layers': sorted(g['layers'])[:4], 'modes': sorted(g['modes'])[:4],
            'hasCfgN': g['hasCfgN'], 'keys': g['keys'],
        })
    games.sort(key=lambda x: (-x['n'], x['title']))

    # ---------- 统计 ----------
    def cnt(field):
        c = {}
        for d in rows:
            v = d.get(field)
            if v: c[v] = c.get(v, 0) + 1
        return sorted(c.items(), key=lambda x: -x[1])

    tier_cnt = {}
    for g in games:
        for t in g['tiers']:
            tier_cnt[t] = tier_cnt.get(t, 0) + 1

    stats = {
        'records': len(rows),
        'games': len(games),
        'playableYes': sum(1 for d in rows if d['playable'] == '是'),
        'playableNo': sum(1 for d in rows if d['playable'] == '否'),
        'playableUnknown': sum(1 for d in rows if d['playable'] not in ('是', '否')),
        'withConfig': sum(1 for d in rows if d['hasCfg']),
        'withNote': sum(1 for d in rows if d['note']),
        'withExe': sum(1 for d in rows if d['exe'].lower().endswith('.exe')),
        'chips': cnt('chip'),
        'chipNames': sorted(set(CHIP[d['chip']] for d in rows if d['chip'])),
        'gpus': cnt('gpu'),
        'layers': cnt('layer'),
        'modes': cnt('mode'),
        'dxvk': cnt('dxvk'),
        'runtimes': cnt('runtime'),
        'fpsTiers': sorted(tier_cnt.items(), key=lambda x: -x[1]),
    }

    out = {
        'builtAt': int(time.time() * 1000),
        'source': os.path.basename(SRC),
        'note': '手机端 PC 游戏实测配置记录（人工测试整理）',
        'chipMap': CHIP,
        'stats': stats,
        'games': games,
        'records': sorted(rows, key=lambda d: (d['game'], d['chipRaw'])),
        'keyIndex': {},     # 运行时由 JS 构建
    }
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False)

    print('=== 已生成', os.path.relpath(OUT), '===')
    print('记录数 %d → 唯一游戏 %d' % (stats['records'], stats['games']))
    print('可玩 是/否/未填 = %d/%d/%d' % (stats['playableYes'], stats['playableNo'], stats['playableUnknown']))
    print('有配置(1) %d · 有备注 %d · 有 exe %d' % (stats['withConfig'], stats['withNote'], stats['withExe']))
    print('机型:', [(CHIP.get(k, k), v) for k, v in stats['chips']])
    print('帧率分档:', stats['fpsTiers'])
    print('文件大小: %.0f KB' % (os.path.getsize(OUT) / 1024))
    print('\nTop 12 游戏:')
    for g in games[:12]:
        print('  %3d条 %-34s 可玩%d 机型%s 档位%s 最佳%s' % (
            g['n'], g['title'][:34], g['okN'], g['chips'], g['tiers'], g['bestLabel'] or '-'))

if __name__ == '__main__':
    main()
